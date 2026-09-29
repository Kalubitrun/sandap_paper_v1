"""Footer with automatic Word page numbers (Phase 5).

Uses real PAGE fields (begin / instrText / end), never typed digits, so
numbers update automatically when the document changes.
"""

from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _field_run(paragraph, instruction: str | None, separator: bool = False):
    run = paragraph.add_run()
    fldChar = run._r.makeelement(qn("w:fldChar"), {})
    if instruction is not None:
        fldChar.set(qn("w:fldCharType"), "begin")
        run._r.append(fldChar)
        instr = paragraph.add_run()
        instr_text = instr._r.makeelement(qn("w:instrText"), {})
        instr_text.set(qn("xml:space"), "preserve")
        instr_text.text = instruction
        instr._r.append(instr_text)
        end = paragraph.add_run()
        endChar = end._r.makeelement(qn("w:fldChar"), {})
        endChar.set(qn("w:fldCharType"), "end")
        end._r.append(endChar)
    else:
        fldChar.set(qn("w:fldCharType"), "separate" if separator else "end")
        run._r.append(fldChar)
    return run


def add_page_number(paragraph, prefix: str = "Page ") -> None:
    paragraph.add_run(prefix)
    _field_run(paragraph, "PAGE")


def render_footer(doc, style) -> None:
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(para)
    for run in para.runs:
        run.font.size = Pt(style.footer_size_pt)
