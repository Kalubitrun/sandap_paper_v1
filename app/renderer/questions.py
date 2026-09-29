"""Questions and sub-questions (Phase 5, refined).

Content preservation rule: headers and sub-questions render the exact
teacher wording/numbering/marks from ``raw_text``. Marks are lifted out
of the line ONLY for positioning (right-aligned tab stop) and the stem
keeps every original character otherwise. Marks are never invented:
without ``marks.raw`` no tab and no mark appear.

Visual hierarchy: the question number run is bold; the stem and marks
runs are regular weight at question size. Plain continuation blocks
(``question.content`` / ``sub.content``) render indented like passages,
merged with sub-questions in original source order.
"""

from __future__ import annotations

import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from app.renderer.options import _cant_split, render_options
from app.renderer.styles import StyleConfig, add_paragraph, add_right_tabstop, borderless

_STRIP_CHARS = " -–—:;,"
_NUMBER_RE = re.compile(
    r"^((?:Q\s*\.?\s*\d+|Question\s+\d+|\d+)\s*[.\)\]:\-]?)\s*(.*?)\s*$",
    re.IGNORECASE | re.DOTALL,
)


def _split_marks(raw_text: str, marks) -> tuple[str, str | None]:
    """Separate the trailing marks annotation for right alignment.

    Mirrors the analyzer's strip so ``stem + marks`` reconstructs the
    original characters exactly (single spaces collapsed).
    """
    if marks and marks.raw and marks.raw in raw_text:
        stem = re.sub(r"\s+", " ", raw_text.replace(marks.raw, "", 1))
        return stem.strip().strip(_STRIP_CHARS), marks.raw
    return raw_text, None


def _split_number(stem: str) -> tuple[str, str]:
    """Split leading question number ("Q6.") from the stem text."""
    m = _NUMBER_RE.match(stem)
    if m and m.group(2):
        return m.group(1).rstrip(), m.group(2)
    if m:
        return m.group(1).rstrip(), ""
    return "", stem


def _header_parts(question) -> tuple[str, str, str | None]:
    raw = (question.raw_text or "").strip()
    if not raw:
        raw = f"{question.raw_number}. {question.text}".strip()
    stem, mark = _split_marks(raw, question.marks)
    number, rest = _split_number(stem)
    return number, rest, mark


def _render_line(doc, number: str, rest: str, mark: str | None,
                 style: StyleConfig, tab_twips: int | None,
                 bold_number: bool, keep_with_next: bool,
                 indent_inches: float = 0.0,
                 space_before_pt: float | None = None):
    para = doc.add_paragraph()
    if indent_inches:
        para.paragraph_format.left_indent = Inches(indent_inches)
    if tab_twips is not None and mark:
        add_right_tabstop(para, tab_twips)
    if number:
        run = para.add_run(number + (" " if rest or mark else ""))
        run.bold = bold_number or None
        run.font.size = Pt(style.question_size_pt)
    if rest:
        run = para.add_run(rest)
        run.font.size = Pt(style.question_size_pt)
    if mark:
        if rest or number:
            para.add_run("\t")
        run = para.add_run(mark)
        run.font.size = Pt(style.question_size_pt)
    para.paragraph_format.space_before = Pt(
        style.question_space_before_pt if space_before_pt is None
        else space_before_pt)
    para.paragraph_format.space_after = Pt(style.question_space_after_pt)
    if keep_with_next:
        para.paragraph_format.keep_with_next = True
    return para


def _sub_kind(label: str) -> str | None:
    """numeric / letter / None for Match-the-Following detection."""
    if label.isdigit():
        return "numeric"
    if len(label) == 1 and label.isalpha():
        return "letter"
    return None


def _is_match_question(question) -> bool:
    """Strict alternating numeric/letter sub-questions (1,a,2,b,...).

    Renderer-side presentation only: anything else (options, passage,
    odd count, fewer than 2 pairs, broken alternation) falls back to
    the normal sub-question layout instead of guessing.
    """
    subs = question.sub_questions
    if (len(subs) < 4 or len(subs) % 2 or question.passage
            or question.options or any(s.options for s in subs)):
        return False
    kinds = [_sub_kind(s.label) for s in subs]
    if any(k is None for k in kinds):
        return False
    return all(k == ("numeric" if i % 2 == 0 else "letter")
               for i, k in enumerate(kinds))


def _render_match_table(doc, question, style: StyleConfig) -> None:
    """Two borderless columns with consistent left edges (never spaces)."""
    subs = question.sub_questions
    table = doc.add_table(rows=len(subs) // 2, cols=2)
    table.autofit = False
    borderless(table)
    widths = (style.match_col_a_inches, style.match_col_b_inches)
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for r in range(len(subs) // 2):
        row = table.rows[r]
        _cant_split(row)
        left = row.cells[0].paragraphs[0]
        run = left.add_run((subs[2 * r].raw_text or "").strip())
        run.font.size = Pt(style.question_size_pt)
        right = row.cells[1].paragraphs[0]
        run = right.add_run((subs[2 * r + 1].raw_text or "").strip())
        run.font.size = Pt(style.question_size_pt)


def _render_content_block(doc, block, style: StyleConfig,
                            keep_with_next: bool = False) -> None:
    """Render one verbatim continuation/passage block, indented like passages."""
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(style.passage_indent_inches)
    run = para.add_run(block.text)
    run.font.size = Pt(style.question_size_pt)
    para.paragraph_format.space_before = Pt(style.question_space_after_pt)
    para.paragraph_format.space_after = Pt(style.question_space_after_pt)
    if keep_with_next:
        para.paragraph_format.keep_with_next = True


def _render_sub(doc, sub, style: StyleConfig, marks_tab_twips: int | None,
                first: bool) -> None:
    raw = (sub.raw_text or "").strip() or f"{sub.raw_label} {sub.text}".strip()
    stem, mark = _split_marks(raw, sub.marks)
    number, rest = _split_number(stem)
    para = _render_line(doc, number, rest, mark, style,
                        marks_tab_twips, bold_number=False,
                        keep_with_next=bool(sub.options or sub.content),
                        indent_inches=0.0, space_before_pt=0.0)
    if first:
        para.paragraph_format.space_before = Pt(
            style.subquestion_first_space_before_pt)
    for block in sub.content:
        _render_content_block(doc, block, style)
    render_options(doc, sub.options, style)


def render_question(doc, question, style: StyleConfig,
                    marks_tab_twips: int | None = None) -> None:
    number, rest, mark = _header_parts(question)
    has_body = bool(question.sub_questions or question.options
                    or question.passage or question.content)
    _render_line(doc, number, rest, mark, style, marks_tab_twips,
                 bold_number=True, keep_with_next=has_body)
    if question.passage:
        passage = doc.add_paragraph()
        passage.paragraph_format.left_indent = Inches(style.passage_indent_inches)
        run = passage.add_run(question.passage)
        run.font.size = Pt(style.question_size_pt)
        passage.paragraph_format.space_before = Pt(style.question_space_after_pt)
        passage.paragraph_format.space_after = Pt(style.question_space_after_pt)
        passage.paragraph_format.keep_with_next = bool(
            question.sub_questions or question.content)
    render_options(doc, question.options, style)
    if _is_match_question(question):
        _render_match_table(doc, question, style)
        return
    # Content blocks and sub-questions in original source order so no
    # teacher paragraph is dropped or reordered.
    stream = [((c.source_blocks + [10 ** 9])[0], "content", c)
              for c in question.content]
    stream += [((s.source_blocks + [10 ** 9])[0], "sub", s)
               for s in question.sub_questions]
    stream.sort(key=lambda item: item[0])
    first_sub = True
    for _, kind, obj in stream:
        if kind == "content":
            _render_content_block(doc, obj, style)
        else:
            _render_sub(doc, obj, style, marks_tab_twips, first_sub)
            first_sub = False


def render_end_marker(doc, end_marker, style: StyleConfig) -> None:
    if not end_marker or not end_marker.get("text"):
        return
    add_paragraph(doc, end_marker["text"].strip(), size_pt=style.question_size_pt,
                  bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  space_before_pt=style.section_space_before_pt)
