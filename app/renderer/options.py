"""MCQ option layout (Phase 5).

Uses the Phase 2 normalized labels, rendered canonically as "(a)".
Four (or two) short options share table rows; long options get their
own lines so nothing overflows the page. Borders are removed.
"""

from __future__ import annotations

from docx.shared import Pt

from app.renderer.styles import StyleConfig, add_paragraph, borderless


def _cant_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(trPr.makeelement(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit",
        {"{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val": "1"}))


def _option_text(option) -> str:
    return f"({option.label}) {option.text}"


def _fits_inline(options, style: StyleConfig) -> bool:
    return (
        len(options) in (2, 4)
        and all(len(o.text) <= style.inline_option_max_chars for o in options)
    )


def render_options(doc, options, style: StyleConfig) -> None:
    if not options:
        return
    if _fits_inline(options, style):
        table = doc.add_table(rows=len(options) // 2, cols=2)
        table.autofit = True
        borderless(table)
        for i in range(0, len(options), 2):
            row = table.rows[i // 2]
            _cant_split(row)
            for col, option in enumerate(options[i:i + 2]):
                cell_para = row.cells[col].paragraphs[0]
                run = cell_para.add_run(_option_text(option))
                run.font.size = Pt(style.option_size_pt)
                cell_para.paragraph_format.space_after = Pt(style.option_space_after_pt)
    else:
        for option in options:
            add_paragraph(doc, _option_text(option), size_pt=style.option_size_pt,
                          space_after_pt=style.option_space_after_pt)
