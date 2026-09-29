"""Renderer refinement regression tests (design pass).

Verifies: invisible layout tables, larger school name, dual logos with
aspect ratio, header/time alignment, right-aligned marks (incl. long
questions), centered sections, page numbers, preserved wording, and the
inline-option fix from the previous bugfix.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.models.document import ParagraphInfo, RawDocument, RunInfo
from app.parser.analyzer import analyze_document, split_inline_options
from app.parser.extractor import clean_text, extract_document
from app.renderer import PaperSettings, SchoolConfig, render_question_paper
from app.validation import validate_paper

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "input" / "amateur_teacher_question_paper.docx"
LOGO = ROOT / "logo.webp"


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    raw = extract_document(INPUT)
    paper = analyze_document(raw)
    result = validate_paper(paper, maximum_marks=50)
    out = tmp_path_factory.mktemp("refine") / "refined.docx"
    render_question_paper(paper, SchoolConfig(), PaperSettings(), result, out)
    return Document(str(out)), paper, result


@pytest.fixture(scope="module")
def doc_xml(built):
    return built


def _tables(doc):
    return doc.tables


# 1. No visible borders on internal tables.
def test_no_visible_table_borders(built):
    doc, _, _ = built
    assert len(_tables(doc)) >= 3
    for table in _tables(doc):
        for m in re.finditer(rb"<w:tblBorders>(.*?)</w:tblBorders>",
                             table._tbl.xml.encode(), re.DOTALL):
            assert b'w:val="single"' not in m.group(1)
            assert b'w:val="nil"' in m.group(1)
        for row in table.rows:
            for cell in row.cells:
                xml = cell._tc.xml
                for tag in ("tcBorders",):
                    assert f"<w:{tag}>" not in xml or 'w:val="nil"' in xml


def _all_paragraphs(doc):
    yield from doc.paragraphs
    for table in _tables(doc):
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


# 2. School name is larger than body text.
def test_school_name_dominant(built):
    doc, _, _ = built
    sizes = [r.font.size.pt for p in _all_paragraphs(doc) for r in p.runs
             if r.text == "VARANASI PUBLIC SCHOOL"]
    assert sizes, "school name run missing"
    body = [r.font.size.pt for p in doc.paragraphs for r in p.runs
            if r.text.startswith("Q12.")]
    assert body and sizes[0] > body[0]
    assert sizes[0] >= 17  # dominant: clearly larger than exam (14pt) and body


# 3. Both logo positions exist when configured.
def test_dual_logos(built):
    doc, _, _ = built
    assert len(doc.inline_shapes) == 2
    widths = [s.width.pt for s in doc.inline_shapes]
    assert widths[0] == pytest.approx(widths[1])


# 4. Logos preserve aspect ratio.
def test_logo_aspect_ratio(built):
    from PIL import Image
    from io import BytesIO
    doc, _, _ = built
    with Image.open(LOGO) as im:
        native = im.size[0] / im.size[1]
    for shape in doc.inline_shapes:
        assert shape.width.pt / shape.height.pt == pytest.approx(native, rel=0.02)


# 5/6. Affiliation top-left, school number top-right.
def test_affiliation_left_school_no_right(built):
    doc, _, _ = built
    top = _tables(doc)[0]
    left, right = top.cell(0, 0).paragraphs[0], top.cell(0, 1).paragraphs[0]
    assert left.text.startswith("Affiliation No. - 2131898")
    assert left.alignment in (None, WD_ALIGN_PARAGRAPH.LEFT)
    assert right.text.startswith("School No. - 71377")
    assert right.alignment == WD_ALIGN_PARAGRAPH.RIGHT


# 7/8. Time left, max marks right (table, never spaces).
def test_time_left_marks_right(built):
    doc, _, _ = built
    row = [t for t in _tables(doc) if "Time:" in t.cell(0, 0).text][0]
    tcell, mcell = row.cell(0, 0).paragraphs[0], row.cell(0, 1).paragraphs[0]
    assert tcell.text.startswith("Time: 3:00 hours")
    assert "  " not in tcell.text
    assert mcell.text == "M.M.: 50"
    assert mcell.alignment == WD_ALIGN_PARAGRAPH.RIGHT


# 9. Question marks are right aligned via tab stop.
def test_marks_right_tabstop(built):
    doc, _, _ = built
    paras = [p for p in doc.paragraphs if p.text.startswith("Q7.")]
    assert len(paras) == 1
    p = paras[0]
    assert 'w:val="right"' in p._p.xml
    assert p.text.endswith("(3)")
    assert "  " not in p.text  # no space-padding for alignment


# 10. Long questions keep marks at the right.
def test_long_question_marks_right(tmp_path):
    long_text = ("Explain how the ideas of liberty, equality and fraternity "
                 "influenced the course of the French Revolution in great detail.")
    raw = RawDocument(
        source="long.docx",
        paragraphs=[
            ParagraphInfo(index=0, original_text="SECTION A",
                          cleaned_text="SECTION A", style="Normal",
                          alignment=None, is_empty=False,
                          runs=[RunInfo(text="SECTION A")]),
            ParagraphInfo(index=1, original_text=f"Q14. {long_text} (4)",
                          cleaned_text=f"Q14. {long_text} (4)",
                          style="Normal", alignment=None, is_empty=False,
                          runs=[RunInfo(text=f"Q14. {long_text} (4)")]),
        ],
        tables=[])
    paper = analyze_document(raw)
    out = tmp_path / "long.docx"
    render_question_paper(paper, SchoolConfig(), PaperSettings(),
                          validate_paper(paper), out)
    d = Document(str(out))
    paras = [p for p in d.paragraphs if p.text.startswith("Q14.")]
    assert len(paras) == 1
    assert 'w:val="right"' in paras[0]._p.xml
    assert paras[0].text.endswith("(4)")
    assert long_text in paras[0].text


# 11. MCQ option tables have no visible borders.
def test_option_tables_borderless(built):
    doc, _, _ = built
    opt_tables = [t for t in _tables(doc) if "(a)" in t.cell(0, 0).text]
    assert opt_tables, "no MCQ option table found"
    for table in opt_tables:
        assert 'w:val="single"' not in table._tbl.xml


# 12. Section headings are centered.
def test_section_headings_centered(built):
    doc, _, _ = built
    heads = [p for p in doc.paragraphs if p.text.startswith("SECTION")]
    assert len(heads) == 2
    assert all(p.alignment == WD_ALIGN_PARAGRAPH.CENTER for p in heads)


# 13. Page numbers exist as fields.
def test_page_numbers(built):
    doc, _, _ = built
    xml = doc.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in xml and "fldChar" in xml


# 14. Original question text remains unchanged.
def test_wording_unchanged(built):
    doc, paper, _ = built
    text_n = _norm("\n".join(
        [p.text for p in doc.paragraphs] +
        [c.text for t in _tables(doc) for r in t.rows for c in r.cells]))
    for section in paper.sections:
        for q in section.questions:
            assert _norm(q.raw_text) in text_n, q.raw_text


# 15. Inline-option regression remains fixed.
def test_inline_option_regression():
    found = split_inline_options(
        "a) First Estate b) Second Estate c) Third Estate d) Nobility")
    assert found is not None and len(found) == 4
    assert [label for label, _, _ in found] == ["a", "b", "c", "d"]
