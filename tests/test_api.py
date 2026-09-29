"""Phase 6 tests: API service layer (no HTTP server needed).

Covers: analyze payload shape, validation/spelling inclusion, friendly
errors (type/size/validity), filename generation, generate round-trip,
validation gate, and settings validation.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from app.api.service import (
    ApiError,
    analyze_document_bytes,
    build_download_filename,
    generate_document_bytes,
)
from app.renderer import PaperSettings

ROOT = Path(__file__).resolve().parents[1]
DOCX = (ROOT / "input" / "amateur_teacher_question_paper.docx").read_bytes()


def settings(**kw) -> PaperSettings:
    base = dict(examination="Half Yearly Examination", session="2026-27",
                class_name="IX", subject="Information Technology",
                subject_code="402", max_marks=50, time="3:00 hours")
    base.update(kw)
    return PaperSettings(**base)


def test_analyze_returns_full_payload():
    out = analyze_document_bytes(DOCX, "paper.docx")
    assert out["summary"] == {"sections": 2, "questions": 19, "instructions": 6,
                              "has_instructions": True, "spelling_suggestions": 0,
                              "has_errors": False}
    assert set(out) == {"filename", "summary", "paper", "validation", "spelling"}
    assert out["validation"]["summary"]["total"] == 17
    assert out["spelling"]["summary"]["total_suggestions"] == 0


def test_analyze_rejects_file_type():
    with pytest.raises(ApiError) as exc:
        analyze_document_bytes(b"hello", "paper.pdf")
    assert exc.value.status_code == 415
    assert ".docx" in exc.value.message


def test_analyze_rejects_empty_and_oversized():
    with pytest.raises(ApiError):
        analyze_document_bytes(b"", "paper.docx")
    with pytest.raises(ApiError) as exc:
        analyze_document_bytes(b"PK\x03\x04" + b"x" * (11 * 1024 * 1024), "p.docx")
    assert exc.value.status_code == 413


def test_analyze_rejects_invalid_docx():
    with pytest.raises(ApiError) as exc:
        analyze_document_bytes(b"not a zip at all....", "paper.docx")
    assert "Unable to read" in exc.value.message


def test_no_traceback_in_errors():
    try:
        analyze_document_bytes(b"junk", "paper.docx")
    except ApiError as exc:
        assert "Traceback" not in exc.message
        assert exc.message


def test_filename_from_settings():
    assert build_download_filename(settings()) == "VPS_Half_Yearly_IX_IT_2026-27.docx"
    assert build_download_filename(settings(examination="PT 1")) == \
        "VPS_PT1_IX_IT_2026-27.docx"


def test_generate_roundtrip_returns_valid_docx():
    blob, name = generate_document_bytes(DOCX, "paper.docx", settings())
    assert name == "VPS_Half_Yearly_IX_IT_2026-27.docx"
    assert blob[:4] == b"PK\x03\x04"
    from docx import Document
    from io import BytesIO
    doc = Document(BytesIO(blob))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(c.text for c in row.cells)
    assert "VARANASI PUBLIC SCHOOL" in "\n".join(parts)


def test_generate_blocked_on_errors():
    from app.models.paper import NormalizedPaper, QuestionInfo, SectionInfo
    from app.api import service as svc
    paper = NormalizedPaper(sections=[SectionInfo(
        title="SECTION A", section_id="A",
        questions=[QuestionInfo(number=6, raw_number="Q6", text="t",
                                raw_text="Q6. t", source_blocks=[1]),
                   QuestionInfo(number=6, raw_number="Q6", text="t2",
                                raw_text="Q6. t2", source_blocks=[2])],
        source_blocks=[0])])
    real_loader = svc._load_paper
    svc._load_paper = lambda data, filename: paper  # noqa: E731
    try:
        with pytest.raises(ApiError) as exc:
            generate_document_bytes(b"PK\x03\x04fake", "p.docx", settings())
        assert exc.value.status_code == 422
        assert "Traceback" not in exc.value.message
    finally:
        svc._load_paper = real_loader


def test_invalid_examination_rejected():
    with pytest.raises(ValueError):
        settings(examination="Unit Test")
