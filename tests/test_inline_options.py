"""Regression tests: single-space inline options (History-paper bug).

Q1 item 4 "a) First Estate b) Second Estate c) Third Estate d) Nobility"
was parsed as ONE option plus a false FEW_OPTIONS warning. The splitter
must handle single-space separators in all supported label spellings.
"""

from __future__ import annotations

from docx import Document

from app.models.document import ParagraphInfo, RawDocument, RunInfo
from app.parser.analyzer import analyze_document, split_inline_options
from app.parser.extractor import clean_text
from app.validation import validate_paper

EXPECTED = [
    ("a", "First Estate"),
    ("b", "Second Estate"),
    ("c", "Third Estate"),
    ("d", "Nobility"),
]


def test_single_space_a_paren():
    assert split_inline_options(
        "a) First Estate b) Second Estate c) Third Estate d) Nobility"
    ) == [("a", "a)", "First Estate"), ("b", "b)", "Second Estate"),
          ("c", "c)", "Third Estate"), ("d", "d)", "Nobility")]


def test_single_space_wrapped_parens():
    assert split_inline_options(
        "(a) First Estate (b) Second Estate (c) Third Estate (d) Nobility"
    ) == [("a", "(a)", "First Estate"), ("b", "(b)", "Second Estate"),
          ("c", "(c)", "Third Estate"), ("d", "(d)", "Nobility")]


def test_single_space_upper_dot():
    assert split_inline_options(
        "A. First Estate B. Second Estate C. Third Estate D. Nobility"
    ) == [("a", "A.", "First Estate"), ("b", "B.", "Second Estate"),
          ("c", "C.", "Third Estate"), ("d", "D.", "Nobility")]


def test_single_space_lower_dot():
    found = split_inline_options(
        "a. First Estate b. Second Estate c. Third Estate d. Nobility")
    assert [(l, t) for l, _, t in found] == EXPECTED


def test_legacy_multispace_format_still_splits():
    found = split_inline_options("a) Email    b) CPU    c) RAM    d) Printer")
    assert [(l, t) for l, _, t in found] == [
        ("a", "Email"), ("b", "CPU"), ("c", "RAM"), ("d", "Printer")]


def test_options_across_paragraphs_unaffected():
    for single in ("a) Windows", "b) Keyboard", "(c) ALU", "D. Printer"):
        assert split_inline_options(single) is None  # single-label path handles these


def test_no_false_split_on_plain_letters():
    # Single option whose text contains ordinary capital letters.
    assert split_inline_options("a) ASDF and JKL;") is None
    # Non-sequential labels: ambiguous, must not guess.
    assert split_inline_options("a) Apples c) Oranges") is None
    # Label not at the start: not an inline run.
    assert split_inline_options("The answer is b) Milk") is None
    # Run not starting at 'a': ambiguous, must not guess.
    assert split_inline_options("b) X c) Y") is None


def _history_paper(texts):
    return RawDocument(
        source="history.docx",
        paragraphs=[ParagraphInfo(
            index=i, original_text=t, cleaned_text=clean_text(t),
            style="Normal", alignment=None, is_empty=(t.strip() == ""),
            runs=[RunInfo(text=t)]) for i, t in enumerate(texts)],
        tables=[])


def test_history_item_end_to_end():
    inline = "a) First Estate b) Second Estate c) Third Estate d) Nobility"
    paper = analyze_document(_history_paper([
        "SECTION A", "Q1. Choose the correct answer.",
        "4. Which estate paid all the taxes in France?", inline,
    ]))
    item = paper.sections[0].questions[0].sub_questions[0]
    assert item.label == "4"
    assert [(o.label, o.text) for o in item.options] == EXPECTED
    # Traceability: complete original paragraph preserved per option.
    assert all(o.source_text == inline for o in item.options)
    assert all(o.source_blocks == [3] for o in item.options)
    # The false warning is gone because detection (not the validator) is fixed.
    result = validate_paper(paper)
    assert result.by_code("FEW_OPTIONS") == []
    assert result.by_code("MCQ_WITHOUT_OPTIONS") == []


def test_renderer_receives_four_separate_options(tmp_path):
    from app.renderer.options import render_options
    from app.renderer.styles import StyleConfig

    inline = "a) First Estate b) Second Estate c) Third Estate d) Nobility"
    paper = analyze_document(_history_paper([
        "SECTION A", "Q1. Choose the correct answer.",
        "4. Which estate paid all the taxes in France?", inline,
    ]))
    options = paper.sections[0].questions[0].sub_questions[0].options
    assert len(options) == 4

    doc = Document()
    render_options(doc, options, StyleConfig())
    assert len(doc.tables) == 1
    cells = [c.text for row in doc.tables[0].rows for c in row.cells]
    assert cells == ["(a) First Estate", "(b) Second Estate",
                     "(c) Third Estate", "(d) Nobility"]
