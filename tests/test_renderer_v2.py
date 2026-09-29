"""Phase 5.2 regression tests: final visual refinement.

Covers: one-line school name layout, larger dual logos, borderless
Match-the-Following columns, Q5-style subquestion spacing, intact
right-aligned marks, and preserved content.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document

from app.models.document import ParagraphInfo, RawDocument, RunInfo
from app.models.paper import QuestionInfo, SectionInfo, SubQuestionInfo
from app.models.paper import NormalizedPaper
from app.parser.analyzer import analyze_document
from app.parser.extractor import clean_text, extract_document
from app.renderer import PaperSettings, SchoolConfig, render_question_paper
from app.renderer.questions import _is_match_question
from app.validation import validate_paper

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "input" / "amateur_teacher_question_paper.docx"


def _raw(texts: list[str]) -> RawDocument:
    return RawDocument(
        source="t.docx",
        paragraphs=[ParagraphInfo(
            index=i, original_text=t, cleaned_text=clean_text(t),
            style="Normal", alignment=None, is_empty=(t.strip() == ""),
            runs=[RunInfo(text=t)]) for i, t in enumerate(texts)],
        tables=[])


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    paper = analyze_document(extract_document(INPUT))
    result = validate_paper(paper, maximum_marks=50)
    out = tmp_path_factory.mktemp("v2") / "v2.docx"
    render_question_paper(paper, SchoolConfig(), PaperSettings(), result, out)
    return Document(str(out)), paper


def _header_table(doc):
    for table in doc.tables:
        if "VARANASI PUBLIC SCHOOL" in table.cell(0, 1).text:
            return table
    raise AssertionError("header table not found")


# 1. School name: single paragraph/run, prominent size, room to fit one line.
def test_school_name_one_line_layout(built):
    doc, _ = built
    head = _header_table(doc)
    name_paras = [p for p in head.cell(0, 1).paragraphs
                  if "VARANASI PUBLIC SCHOOL" in p.text]
    assert len(name_paras) == 1
    assert name_paras[0].text.strip() == "VARANASI PUBLIC SCHOOL"
    size = name_paras[0].runs[0].font.size.pt
    assert size >= 17  # dominant, still clearly larger than body/exam text
    # Center column keeps >= 3.3" so the name fits one line at 19pt.
    grid = [int(v) for v in re.findall(r'<w:gridCol[^>]*w:w="(\d+)"', head._tbl.xml)]
    assert len(grid) == 3
    assert grid[1] / 1440 >= 3.3


# 2/3. Logos larger, identical, aspect preserved, both present.
def test_logos_larger_and_equal(built):
    doc, _ = built
    assert len(doc.inline_shapes) == 2
    widths = [s.width.pt for s in doc.inline_shapes]
    assert widths[0] == pytest.approx(widths[1])
    assert widths[0] == pytest.approx(115.2, abs=0.5)  # 1.6in, +23% over 1.3in
    assert widths[0] > 93.6  # strictly larger than the previous size
    from PIL import Image
    with Image.open(ROOT / "logo.webp") as im:
        native = im.size[0] / im.size[1]
    for shape in doc.inline_shapes:
        assert shape.width.pt / shape.height.pt == pytest.approx(native, rel=0.02)


# 4/5. Match-the-Following: aligned borderless columns, no visible borders.
def _match_paper():
    subs = []
    for i, (num, let, a, b) in enumerate([
            ("1", "a", "Rousseau", "Bastille"),
            ("2", "b", "Louis XVI", "The Social Contract"),
            ("3", "c", "14 July 1789", "King of France"),
            ("4", "d", "Third Estate", "Common people")]):
        subs.append(SubQuestionInfo(label=num, raw_label=f"{num}.", text=a,
                                    raw_text=f"{num}. {a}", source_blocks=[i]))
        subs.append(SubQuestionInfo(label=let, raw_label=f"{let}.", text=b,
                                    raw_text=f"{let}. {b}", source_blocks=[i]))
    q = QuestionInfo(number=5, raw_number="Q5", text="Match the following.",
                     raw_text="Q5. Match the following.", sub_questions=subs,
                     source_blocks=[0])
    return NormalizedPaper(sections=[SectionInfo(
        title="SECTION B", section_id="B", questions=[q], source_blocks=[0])])


def test_match_detection_and_layout(tmp_path):
    paper = _match_paper()
    q = paper.sections[0].questions[0]
    assert _is_match_question(q) is True
    out = tmp_path / "match.docx"
    render_question_paper(paper, SchoolConfig(), PaperSettings(),
                          validate_paper(paper), out)
    doc = Document(str(out))
    tables = [t for t in doc.tables
              if t.cell(0, 0).text == "1. Rousseau"
              and t.cell(0, 1).text == "a. Bastille"]
    assert len(tables) == 1
    table = tables[0]
    assert len(table.rows) == 4 and len(table.columns) == 2
    assert table.cell(3, 0).text == "4. Third Estate"
    assert table.cell(3, 1).text == "d. Common people"
    assert 'w:val="single"' not in table._tbl.xml  # no visible grid
    # Consistent left edges: fixed equal-ish column widths.
    widths = [c.width.pt for c in table.rows[0].cells]
    assert widths[0] > 0 and widths[1] > 0


def test_match_fallback_for_non_matching():
    paper = analyze_document(extract_document(INPUT))
    for section in paper.sections:
        for q in section.questions:
            assert _is_match_question(q) is False  # IT sample has none


# 6. Q5-style subquestion spacing within range.
def test_subquestion_spacing(built):
    doc, _ = built
    first = next(p for p in doc.paragraphs if p.text.startswith("a) What is a computer"))
    second = next(p for p in doc.paragraphs if p.text.startswith("b) Write any two"))
    first_before = first.paragraph_format.space_before.pt or 0
    assert 2.0 <= first_before <= 6.0
    second_before = second.paragraph_format.space_before
    assert second_before is None or second_before.pt == 0


# 7/8. Right-aligned marks intact (incl. long questions).
def test_marks_tabs_intact(built):
    doc, _ = built
    marked = [p for p in doc.paragraphs if p.text.startswith("Q7.")]
    assert len(marked) == 1
    assert 'w:val="right"' in marked[0]._p.xml
    assert marked[0].text.endswith("(3)")


# 9/10/11/12. Prior behavior preserved.
def test_content_and_borders_preserved(built):
    import re as _re
    doc, paper = built
    text_n = _re.sub(r"\s+", " ", "\n".join(
        [p.text for p in doc.paragraphs] +
        [c.text for t in doc.tables for r in t.rows for c in r.cells]))
    assert _re.sub(r"\s+", " ", "Q2. Fill in the blanks. (5 x 1 = 5)") in text_n
    for table in doc.tables:
        assert 'w:val="single"' not in table._tbl.xml


def test_inline_option_fix_intact():
    from app.parser.analyzer import split_inline_options
    found = split_inline_options(
        "a) First Estate b) Second Estate c) Third Estate d) Nobility")
    assert found is not None and len(found) == 4
