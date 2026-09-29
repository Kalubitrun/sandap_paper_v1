"""Phase 1 tests: DOCX -> raw representation.

Verifies: document loads, order preserved, original text untouched,
blanks distinguishable, tables detected, run formatting readable.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.parser.extractor import clean_text, extract_document

DOCX = Path(__file__).resolve().parents[1] / "input" / "amateur_teacher_question_paper.docx"


def test_document_loads_successfully():
    raw = extract_document(DOCX)
    assert raw.source == DOCX.name
    assert len(raw.paragraphs) == 91
    assert isinstance(raw.tables, list)


def test_paragraphs_extracted_in_correct_order():
    raw = extract_document(DOCX)
    assert [p.index for p in raw.paragraphs] == list(range(len(raw.paragraphs)))
    assert raw.paragraphs[0].original_text == "VARANASI PUBLIC SCHOOL"
    assert raw.paragraphs[1].original_text == "HALF YEARLY EXAMINATION 2026-27"
    assert raw.paragraphs[90].original_text.strip() == "END"


def test_original_text_is_preserved_exactly():
    # Compare against python-docx ground truth: extractor must not alter text.
    doc = DocxDocument(str(DOCX))
    raw = extract_document(DOCX)
    assert len(raw.paragraphs) == len(doc.paragraphs)
    for para, info in zip(doc.paragraphs, raw.paragraphs):
        assert info.original_text == para.text
    # Spot-check a tricky paragraph: inline 4-space option separator kept.
    q4_opts = raw.paragraphs[35]
    assert q4_opts.original_text == "a) Email    b) CPU    c) RAM    d) Printer"
    # cleaned_text is derived, original untouched.
    assert q4_opts.cleaned_text == "a) Email b) CPU c) RAM d) Printer"
    assert q4_opts.original_text != q4_opts.cleaned_text
    # Leading spaces preserved in original, stripped in cleaned.
    end = raw.paragraphs[90]
    assert end.original_text.startswith(" ")
    assert end.cleaned_text == "END"


def test_blank_paragraphs_are_distinguishable():
    raw = extract_document(DOCX)
    blanks = [p.index for p in raw.paragraphs if p.is_empty]
    assert blanks == [7, 15, 17, 41, 48, 54, 89]
    for p in raw.paragraphs:
        assert p.is_empty == (p.original_text.strip() == "")
    non_blank = raw.paragraphs[0]
    assert non_blank.is_empty is False


def test_tables_are_detected():
    # This fixture document contains no tables -> empty list (not None).
    raw = extract_document(DOCX)
    assert raw.tables == []


def test_tables_extracted_from_synthetic_doc(tmp_path):
    doc = DocxDocument()
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A1"
    t.cell(1, 1).text = "B2"
    p = tmp_path / "with_table.docx"
    doc.save(str(p))
    raw = extract_document(p)
    assert len(raw.tables) == 1
    assert raw.tables[0].rows == 2
    assert raw.tables[0].cols == 2
    texts = {(c.row, c.col): c.text for c in raw.tables[0].cells}
    assert texts[(0, 0)] == "A1"
    assert texts[(1, 1)] == "B2"


def test_basic_run_formatting_can_be_read():
    raw = extract_document(DOCX)
    # Bold school name with explicit 15pt size.
    first = raw.paragraphs[0]
    assert len(first.runs) == 1
    assert first.runs[0].bold is True
    assert first.runs[0].font_size_pt == pytest.approx(15.0)
    # A body line is explicitly non-bold.
    assert raw.paragraphs[1].runs[0].bold is False
    # Every paragraph exposes run list; style + alignment fields exist.
    for p in raw.paragraphs:
        assert p.style == "Normal"
        assert p.alignment is None or isinstance(p.alignment, str)
        assert isinstance(p.runs, list) and len(p.runs) >= 1


def test_clean_text_helper_never_mutates_input():
    s = "   hello    world   "
    assert clean_text(s) == "hello world"
    assert s == "   hello    world   "


def test_missing_file_raises():
    with pytest.raises(FileNotFoundError):
        extract_document("input/does_not_exist.docx")
