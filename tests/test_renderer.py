"""Phase 5 tests: VPS DOCX renderer (Phase 5).

Verifies header/school info, paper settings, content preservation,
numbering/marks integrity, page numbers, layout mechanisms, the
validation gate, and the full pipeline round-trip.
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app.parser.analyzer import analyze_document
from app.parser.extractor import extract_document
from app.renderer import (
    PaperSettings,
    RenderBlockedError,
    SchoolConfig,
    render_question_paper,
)
from app.validation import validate_paper

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "input" / "amateur_teacher_question_paper.docx"
LOGO = ROOT / "logo.webp"


def _pipeline():
    raw = extract_document(INPUT)
    paper = analyze_document(raw)
    result = validate_paper(paper, maximum_marks=50)
    return paper, result


@pytest.fixture(scope="module")
def rendered(tmp_path_factory):
    paper, result = _pipeline()
    out = tmp_path_factory.mktemp("vps") / "vps_question_paper.docx"
    render_question_paper(paper, SchoolConfig(), PaperSettings(), result, out)
    return out, paper, result


@pytest.fixture(scope="module")
def doc(rendered):
    return Document(str(rendered[0]))


def _full_text(doc) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


# 1. DOCX is generated
def test_docx_generated(rendered):
    out, _, _ = rendered
    assert out.is_file() and out.stat().st_size > 0
    Document(str(out))  # opens cleanly


# 2-5. fixed school information
def test_school_name(doc):
    assert "VARANASI PUBLIC SCHOOL" in _full_text(doc)


def test_address(doc):
    assert "BANGALIPUR, RAJATALAB, VARANASI" in _full_text(doc)


def test_affiliation_number(doc):
    assert "2131898" in _full_text(doc)


def test_school_number(doc):
    assert "71377" in _full_text(doc)


# 6. phone number comes from configuration (never invented)
def test_phone_from_config_only(rendered, tmp_path_factory):
    _, paper, result = rendered
    default_text = _full_text(Document(str(rendered[0])))
    assert "Phone" not in default_text  # default config has no phone: omitted

    out2 = tmp_path_factory.mktemp("vps2") / "with_phone.docx"
    render_question_paper(paper, SchoolConfig(phone_number="+91-98765 43210"),
                          PaperSettings(), result, out2)
    assert "+91-98765 43210" in _full_text(Document(str(out2)))


# 7. existing project logo is used on BOTH sides (same image, converted to PNG)
def test_project_logo_used(rendered):
    out, _, _ = rendered
    assert LOGO.is_file()
    doc = Document(str(out))
    assert len(doc.inline_shapes) == 2
    with zipfile.ZipFile(out) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        assert len(media) == 1 and media[0].endswith(".png")
        embedded = z.read(media[0])
    from PIL import Image
    expected = BytesIO()
    with Image.open(LOGO) as im:
        im.save(expected, format="PNG")
    assert embedded == expected.getvalue()


# 8-14. dynamic paper settings
def test_examination_name(doc):
    assert "Half Yearly Examination" in _full_text(doc)


def test_session(doc):
    assert "2026-27" in _full_text(doc)


def test_class(doc):
    assert "CLASS- IX" in _full_text(doc)


def test_subject(doc):
    assert "Information Technology" in _full_text(doc)


def test_subject_code(doc):
    assert "(402)" in _full_text(doc)


def test_maximum_marks(doc):
    assert "M.M.: 50" in _full_text(doc)


def test_time(doc):
    assert "Time: 3:00 hours" in _full_text(doc)


# 15-19. teacher content blocks
def test_general_instructions(doc):
    text = _full_text(doc)
    for line in ("General Instructions:",
                 "a) Please read all the questions carefully.",
                 "f) Marks are given against each question."):
        assert line in text


def test_sections(doc):
    text = _full_text(doc)
    assert "SECTION A OBJECTIVE TYPE QUESTIONS 24 MARKS" in text
    assert "SECTION B SUBJECTIVE TYPE QUESTIONS (26 MARKS)" in text


def test_questions(doc):
    import re
    text = re.sub(r"\s+", " ", _full_text(doc))
    assert "Q1. Choose the correct answer." in text
    assert "Q2. Fill in the blanks. (5 x 1 = 5)" in text
    assert "Q21. Case Study" in text


def test_options(doc):
    text = _full_text(doc)
    for opt in ("(a) Windows", "(b) Keyboard", "(c) ALU", "(d) Printer",
                "(a) Words Per Machine"):
        assert opt in text


def test_subquestions(doc):
    text = _full_text(doc)
    assert "i) CPU" in text and "iii) ROM" in text
    assert "1) What is BPO?" in text
    assert "a) What is a computer?" in text


# 20. original teacher wording preserved (whitespace-normalized: the
# renderer uses a tab stop to right-align marks, never spaces)
def test_teacher_wording_preserved(rendered):
    import re
    _, paper, _ = rendered
    text = _full_text(Document(str(rendered[0])))
    norm = lambda s: re.sub(r"\s+", " ", s).strip()
    text_n = norm(text)
    checked = 0
    for section in paper.sections:
        assert section.title in text
        for q in section.questions:
            assert norm(q.raw_text) in text_n, q.raw_text
            checked += 1
            if q.passage:
                assert q.passage in text
            for opt in q.options:
                assert opt.text in text
            for sq in q.sub_questions:
                assert norm(sq.raw_text) in text_n, sq.raw_text
                checked += 1
                for opt in sq.options:
                    assert opt.text in text
    for inst in paper.instructions[1:]:
        assert inst.text in text
    assert checked > 30


# 21. question numbering not changed
def test_numbering_unchanged(doc):
    text = _full_text(doc)
    assert "Q4." not in text and "Q5." not in text
    assert text.index("Q3.") < text.index("Q6.")


# 22. missing marks not invented
def test_missing_marks_not_invented(doc):
    paras = [p.text for p in doc.paragraphs]
    q12 = next(p for p in paras if p.startswith("Q12."))
    assert q12 == "Q12. Explain the different types of computer memory."
    q8 = next(p for p in paras if p.startswith("Q8."))
    assert q8 == "Q8. Answer the following:"
    assert "(3)" in next(p for p in paras if p.startswith("Q7."))


# 23. automatic page numbers (field, not typed digits)
def test_page_number_field(doc):
    xml = doc.sections[0].footer.paragraphs[0]._p.xml
    assert "PAGE" in xml and "fldChar" in xml


# layout: tables/tabs, never spaces, for alignment
def test_no_space_alignment(doc):
    for p in doc.paragraphs:
        assert "  " not in p.text, f"double spaces in {p.text!r}"
    assert any(len(t.rows[0].cells) == 2 for t in doc.tables)


def test_a4_page_setup(doc):
    section = doc.sections[0]
    assert section.page_width.pt == pytest.approx(595.45, abs=1)
    assert section.page_height.pt == pytest.approx(841.7, abs=1)


# validation gate + config model
def test_error_gate_blocks_generation(tmp_path):
    from app.models.paper import NormalizedPaper, QuestionInfo, SectionInfo
    paper = NormalizedPaper(sections=[SectionInfo(
        title="SECTION A", section_id="A",
        questions=[QuestionInfo(number=6, raw_number="Q6", text="t",
                                raw_text="Q6. t", source_blocks=[1]),
                   QuestionInfo(number=6, raw_number="Q6", text="t2",
                                raw_text="Q6. t2", source_blocks=[2])],
        source_blocks=[0])])
    result = validate_paper(paper)
    assert result.has_errors
    with pytest.raises(RenderBlockedError):
        render_question_paper(paper, SchoolConfig(), PaperSettings(), result,
                              tmp_path / "blocked.docx")
    assert not (tmp_path / "blocked.docx").exists()


def test_invalid_examination_rejected():
    with pytest.raises(ValueError):
        PaperSettings(examination="Midterm Madness")


# 24. full pipeline round-trip integrity
def test_pipeline_roundtrip_preserves_counts(rendered):
    _, paper, result = rendered
    raw = extract_document(INPUT)
    content_idx = {p.index for p in raw.paragraphs if not p.is_empty}
    covered = {m.source_block for m in paper.metadata}
    covered |= {i.source_blocks[0] for i in paper.instructions if i.source_blocks}
    for section in paper.sections:
        covered.update(section.source_blocks)
        for q in section.questions:
            covered.update(q.source_blocks)
    if paper.end_marker:
        covered.add(paper.end_marker["source_block"])
    assert covered == content_idx  # every content para accounted for
    assert not result.has_errors
