"""Content-preservation regression tests (passage/extract bug).

No English source DOCX exists in the repo, so the fixture below is a
reconstruction from the reported structures: Q1 reading passage and Q6
literature extract, with the task's own example wording. It exercises
the generic mechanism (plain paragraphs -> ContentBlock -> rendered in
source order), not any question-specific special case.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document

from app.models.document import ParagraphInfo, RawDocument, RunInfo
from app.parser.analyzer import analyze_document
from app.parser.extractor import clean_text, extract_document
from app.renderer import PaperSettings, SchoolConfig, render_question_paper
from app.validation import validate_paper

IT_DOCX = Path(__file__).resolve().parents[1] / "input" / "amateur_teacher_question_paper.docx"

PASSAGE = ("Books have always played an important role in human life. "
           "They give us knowledge, improve our imagination and help us "
           "understand people and places that we may never visit.")
EXTRACT = ("The little bird looked at the open sky and wondered whether "
           "it could ever fly beyond the tall trees.")
TRAILING = "Answer all parts in complete sentences."

ENGLISH_TEXTS = [
    "SECTION A",
    "Q1. Read the following passage and answer the questions that follow. (5 x 1 = 5)",
    PASSAGE,
    "a) What do books give us?",
    "b) How do books help our imagination?",
    "c) What do books help us understand?",
    "Q6. Read the following extract and answer the questions. (5)",
    EXTRACT,
    "a) What did the bird look at?",
    "b) What was the bird afraid of?",
    TRAILING,
]


def english_raw() -> RawDocument:
    return RawDocument(
        source="english.docx",
        paragraphs=[ParagraphInfo(
            index=i, original_text=t, cleaned_text=clean_text(t),
            style="Normal", alignment=None, is_empty=(t.strip() == ""),
            runs=[RunInfo(text=t)]) for i, t in enumerate(ENGLISH_TEXTS)],
        tables=[])


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def _full(doc) -> str:
    """Body text in document order (paragraphs and tables interleaved)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    parts: list[str] = []
    for child in doc.element.body:
        if child.tag.endswith("}p"):
            parts.append(Paragraph(child, doc).text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, doc).rows:
                parts.extend(c.text for c in row.cells)
    return _norm("\n".join(parts))


def _sequence(paper) -> list[str]:
    """Every content string in source order."""
    seq: list[str] = []
    for inst in paper.instructions:
        seq.append(inst.text)
    for section in paper.sections:
        if section.title:
            seq.append(section.title)
        for q in section.questions:
            seq.append(q.raw_text)
            if q.passage:
                seq.append(q.passage)
            for o in q.options:
                seq.append(o.text)
            items: list[tuple[int, str]] = []
            for c in q.content:
                items.append((min(c.source_blocks), c.text))
            for sq in q.sub_questions:
                items.append((min(sq.source_blocks), sq.raw_text))
                for c in sq.content:
                    items.append((min(c.source_blocks), c.text))
                for o in sq.options:
                    items.append((min(o.source_blocks), o.text))
            items.sort(key=lambda item: item[0])
            seq.extend(t for _, t in items)
    if paper.end_marker and paper.end_marker.get("text"):
        seq.append(paper.end_marker["text"].strip())
    return [_norm(s) for s in seq if _norm(s)]


@pytest.fixture(scope="module")
def english(tmp_path_factory):
    paper = analyze_document(english_raw())
    out = tmp_path_factory.mktemp("content") / "english.docx"
    render_question_paper(paper, SchoolConfig(), PaperSettings(),
                          validate_paper(paper), out)
    return paper, _full(Document(str(out)))


# 1/2. Passage and extract stored structurally and rendered.
def test_reading_passage_appears(english):
    paper, text = english
    q1 = paper.sections[0].questions[0]
    assert len(q1.content) == 1 and q1.content[0].text == PASSAGE
    assert q1.content[0].source_blocks == [2]
    # Stored structurally, not merged into the stem (merging hid the drop).
    assert PASSAGE not in q1.text
    assert PASSAGE in text


def test_literature_extract_appears(english):
    paper, text = english
    q6 = paper.sections[0].questions[1]
    assert q6.content[0].text == EXTRACT
    assert q6.content[0].source_blocks == [7]
    assert EXTRACT in text


# 3. Case-study passage still appears (IT sample).
def test_case_study_still_appears(tmp_path):
    paper = analyze_document(extract_document(IT_DOCX))
    q21 = [q for s in paper.sections for q in s.questions if q.number == 21][0]
    assert q21.passage is not None
    out = tmp_path / "it.docx"
    render_question_paper(paper, SchoolConfig(), PaperSettings(),
                          validate_paper(paper), out)
    assert q21.passage in _full(Document(str(out)))


# 4. Trailing continuation paragraphs still appear, in order.
def test_continuation_appears_in_order(english):
    paper, text = english
    q6 = paper.sections[0].questions[1]
    assert q6.content[-1].text == TRAILING
    assert text.index(_norm(EXTRACT)) \
        < text.index("What did the bird look at?") \
        < text.index(_norm(TRAILING))


# 5/6. Wording unchanged, paragraph order unchanged.
def test_wording_and_order(english):
    _, text = english
    cursor = 0
    for expected in [PASSAGE,
                     "a) What do books give us?",
                     "Q6. Read the following extract and answer the questions. (5)",
                     EXTRACT,
                     "a) What did the bird look at?",
                     TRAILING]:
        found = text.find(_norm(expected), cursor)
        assert found >= cursor, expected
        cursor = found + 1


# 7. Source blocks remain traceable.
def test_source_blocks_traceable(english):
    paper, _ = english
    q1 = paper.sections[0].questions[0]
    assert 2 in q1.source_blocks  # passage block attributed to Q1
    assert q1.sub_questions[0].source_blocks == [3]
    by_idx = {p.index: p for p in english_raw().paragraphs}
    assert by_idx[2].original_text == PASSAGE


# 8 + CRITICAL. Round-trip: every normalized content string present, in order.
def test_round_trip_no_content_dropped(english):
    paper, text = english
    cursor = 0
    for expected in _sequence(paper):
        found = text.find(expected, cursor)
        assert found >= cursor, f"dropped or reordered: {expected[:60]!r}"
        cursor = found + len(expected)


def test_it_sample_round_trip(tmp_path):
    paper = analyze_document(extract_document(IT_DOCX))
    out = tmp_path / "it.docx"
    render_question_paper(paper, SchoolConfig(), PaperSettings(),
                          validate_paper(paper), out)
    text = _full(Document(str(out)))
    cursor = 0
    for expected in _sequence(paper):
        found = text.find(expected, cursor)
        assert found >= cursor, f"dropped or reordered: {expected[:60]!r}"
        cursor = found + len(expected)


def test_validator_accepts_content_only_stem():
    paper = analyze_document(english_raw())
    result = validate_paper(paper)
    assert result.by_code("EMPTY_QUESTION") == []
